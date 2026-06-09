from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "УП03_Этап4" / "docs" / "QUALITY_REPORT.docx"


def cell_text(cell, text, bold=False):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell_text(table.rows[0].cells[index], header, True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell_text(cells[index], value)
    document.add_paragraph()


def heading(document, text):
    paragraph = document.add_heading(text, level=1)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(31, 78, 121)


def paragraph(document, text):
    item = document.add_paragraph(text)
    item.paragraph_format.space_after = Pt(6)
    item.paragraph_format.line_spacing = 1.15


def bullet(document, text):
    document.add_paragraph(text, style="List Bullet")


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
document.styles["Normal"].font.name = "Arial"
document.styles["Normal"].font.size = Pt(10.5)

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run(
    "Отчёт по качеству программной системы\n"
    "УП.03. Этап 4. Тестирование и диагностика качества"
)
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(31, 78, 121)

heading(document, "1. Объект и дата проверки")
paragraph(
    document,
    "Проверка выполнена 9 июня 2026 года. Репозиторий: "
    "https://github.com/Desai0/site. Production-сервис: "
    "https://desaichk.com. Страница заявок: "
    "https://desaichk.com/track-requests.html.",
)

heading(document, "2. Выполненные проверки")
add_table(
    document,
    ["Проверка", "Фактический результат", "Статус"],
    [
        ["Главная страница", "HTTP 200, 27 547 байт, 0.201 с", "Пройдено"],
        ["Страница заявок", "HTTP 200, 9 845 байт, 0.234 с", "Пройдено"],
        ["Успешный POST", "HTTP 201, создана запись id=6", "Пройдено"],
        ["GET созданной записи", "HTTP 200, запись id=6 получена", "Пройдено"],
        ["Ошибочный GET", "HTTP 404 с понятным JSON-сообщением", "Пройдено"],
        ["Ошибочный POST", "HTTP 400 с перечнем обязательных полей", "Пройдено"],
        ["Production build", "Vite 7.3.2, 26 модулей, 373 мс", "Пройдено"],
        ["Backend syntax", "Два node --check, exit code 0", "Пройдено"],
        ["Docker", "Оба production-контейнера находятся в состоянии Up", "Пройдено"],
        ["Lighthouse", "75 / 95 / 100 / 83; LCP 51.6 с", "С замечаниями"],
    ],
)

heading(document, "3. Производительность")
paragraph(
    document,
    "Lighthouse показал Performance 75, Accessibility 95, Best Practices 100 "
    "и SEO 83. FCP составил 1.1 с, Speed Index 1.2 с, Total Blocking Time 0 мс, "
    "CLS 0. Основная проблема — LCP 51.6 с при переданном объёме 10 118 KiB.",
)
paragraph(
    document,
    "Дополнительно выполнено 10 GET-запросов к главной странице: все запросы "
    "успешны, среднее время ответа 0.2264 с, минимум 0.1733 с, максимум 0.3103 с.",
)

heading(document, "4. Дефекты")
add_table(
    document,
    ["ID", "Дефект", "Критичность", "Статус"],
    [
        ["BUG-01", "favicon.ico и favicon.svg возвращают HTTP 404", "Низкая", "Открыт"],
        ["BUG-02", "LCP 51.6 с из-за общего payload около 10 МБ", "Высокая", "Открыт"],
        ["BUG-03", "Повреждённый JSON может вернуть общий HTTP 500", "Средняя", "Запланирован"],
        ["BUG-04", "JSON-хранилище не защищено от конкурентной записи", "Высокая", "Запланирован"],
        ["BUG-05", "Scraper зависит от HTML внешних providers", "Средняя", "Снижен"],
    ],
)

heading(document, "5. Риски и меры снижения")
for text in [
    "Недоступность внешнего provider: использовать timeout, fallback и Redis-кэш.",
    "Потеря данных: перейти с JSON/SQLite-прототипа на серверную БД с транзакциями.",
    "Несовместимость SpaceTimeDB bindings: применять одинаковые версии CLI и SDK.",
    "Рост realtime-нагрузки: сохранять throttling, batching и TTL.",
    "Ошибка Docker/nginx после обновления: использовать healthcheck и документированный deploy.",
]:
    bullet(document, text)

heading(document, "6. Логи и запуск")
paragraph(
    document,
    "Контейнеры site-frontend-prod и track-backend-prod находятся в состоянии Up. "
    "Vite запустился за 122–222 мс, backend слушает 0.0.0.0:18080. В полученном "
    "фрагменте Docker-логов критические ошибки запуска отсутствуют.",
)

heading(document, "7. Вывод")
paragraph(
    document,
    "Проект работоспособен и пригоден для учебной демонстрации. Все основные "
    "UI/API-сценарии и ошибочные запросы обработаны ожидаемо. Перед дальнейшей "
    "эксплуатацией необходимо оптимизировать крупное изображение, исправить favicon, "
    "перенести хранилище в полноценную БД и улучшить обработку некорректного JSON. "
    "Для повторных scraper-запросов целесообразно использовать Redis-кэш.",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
