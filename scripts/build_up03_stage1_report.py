from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, RGBColor, Inches


OUTPUT = "01_Входной_аудит_проекта.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(3)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(header_cells[index], header, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)

    document.add_paragraph()
    return table


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(31, 78, 121)


def add_paragraph(document, text):
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Отчёт по учебной практике УП.03\nЭтап 1. Входной аудит проекта и подготовка к запуску")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Специальность: 09.02.07 Информационные системы и программирование\n")
    subtitle.add_run("Профессиональный модуль: ПМ.03. Сопровождение и обслуживание программного обеспечения компьютерных систем\n")
    subtitle.add_run("Объём этапа: 12 часов")

    add_heading(document, "1. Паспорт проекта")
    add_table(
        document,
        ["Поле", "Содержание"],
        [
            ["Название проекта", "Desaichk Portfolio / интерактивный сайт с realtime-курсорами и сервисом заявок на треки"],
            [
                "Что делает программа",
                "Проект представляет собой веб-сайт с интерактивной страницей портфолио, синхронизацией курсоров пользователей через SpaceTimeDB, сохранением временной истории рисунков и отдельной страницей для отправки заявок на поиск ссылок скачивания музыкальных треков.",
            ],
            [
                "Для кого программа",
                "Посетители сайта, владелец портфолио, пользователи страницы заявок, администратор проекта.",
            ],
            [
                "Связь с УП.02",
                "Проект является продолжением собственного проекта, начатого на УП.02. На УП.02 были подготовлены идея, прототип интерфейса и базовая реализация. На УП.03 проект рассматривается как сопровождаемая программная система: проверяются состав, запуск, зависимости, конфигурация, документация и готовность к дальнейшей эксплуатации.",
            ],
            [
                "Стек",
                "HTML, CSS, JavaScript, TypeScript, Vite, SpaceTimeDB, Node.js, Docker, Nginx. Для backend scraper используется Node.js HTTP-сервер и файловое JSON-хранилище как временный вариант.",
            ],
            ["Ссылка на репозиторий", "https://github.com/Desai0/site или актуальная ссылка на GitHub-репозиторий проекта"],
            ["Текущее состояние", "Проект запускается частично/полностью: frontend собирается через Vite, realtime-функции работают через SpaceTimeDB, backend scraper протестирован через HTTP-запросы."],
            [
                "Главная проблема на входе",
                "Проект состоит из нескольких частей, поэтому требуется актуализировать инструкцию запуска, .env.example, описание backend scraper и порядок деплоя на VPS.",
            ],
        ],
    )

    add_heading(document, "2. Что было сделано в УП.02")
    add_paragraph(
        document,
        "В рамках УП.02 был начат собственный проект веб-сайта. Были подготовлены прототип интерфейса, структура страниц, базовая логика взаимодействия с пользователем и первичная документация. Проект изначально рассматривался как личный веб-интерфейс, но в процессе развития получил дополнительные серверные функции."
    )
    add_paragraph(
        document,
        "На этапе УП.03 работа продолжается не как создание нового проекта с нуля, а как входной аудит и подготовка существующей разработки к сопровождению. Дополнительно намечено развитие серверной части: scraper будет взаимодействовать с базой данных на отдельном сервере, а старый подход к локальному хранению данных планируется заменить на более надёжное решение. В дальнейшем возможно использование полноценной БД вместо временного SQLite/файлового хранения, а также добавление Redis для кэширования часто запрашиваемых результатов."
    )

    add_heading(document, "3. Проверка состава проекта")
    add_table(
        document,
        ["Файл/папка", "Что есть сейчас", "Назначение", "Что сделать"],
        [
            ["README.md", "Есть", "Краткое описание проекта, стек, запуск, связь с SpaceTimeDB", "Актуализировать раздел backend scraper и Docker-деплой"],
            ["package.json", "Есть", "Скрипты frontend-сборки, Vite, генерация bindings SpaceTimeDB", "ОК"],
            [".env.production", "Есть", "Production-настройки SpaceTimeDB host/db name", "Подготовить .env.example без секретов"],
            ["index.html", "Есть", "Основная страница сайта", "ОК"],
            ["track-requests.html", "Есть", "Страница отправки заявки на поиск трека", "ОК"],
            ["script.js", "Есть", "UI-логика, локальное рисование, декоративные эффекты", "ОК"],
            ["style.css", "Есть", "Основные стили сайта и курсоров", "ОК"],
            ["src/cursor-sync.ts", "Есть", "Realtime-синхронизация курсоров и история рисунков через SpaceTimeDB", "Убрать лишние debug-логи перед финальной сдачей"],
            ["src/track-page.ts", "Есть", "Frontend-логика страницы заявок", "Можно дополнить отображением результата сразу после submit"],
            ["src/module_bindings/", "Есть", "Сгенерированные bindings SpaceTimeDB", "Не редактировать вручную"],
            ["spacetimedb/spacetimedb/src/index.ts", "Есть", "Серверный модуль SpaceTimeDB: cursor, stroke history, reducers/procedure", "ОК, при изменениях делать publish/generate"],
            ["backend_scraper/", "Есть", "Node.js backend для заявок и поиска downloadUrl", "Добавить описание в README и production-инструкцию"],
            ["backend_scraper/data/", "Игнорируется git", "Локальное хранилище заявок", "В дальнейшем заменить на нормальную БД"],
            ["dist/", "Генерируется", "Production-сборка сайта", "Не хранить в git"],
            ["screenshots/", "Нужно подготовить для сдачи", "Скриншоты структуры, запуска и работы", "Добавить фактические скриншоты"],
            ["logs/", "Нужно подготовить для сдачи", "Логи запуска frontend/backend", "Добавить app_start_log.txt"],
        ],
    )

    add_heading(document, "4. Зависимости и версии")
    add_table(
        document,
        ["Зависимость", "Версия/требование", "Команда проверки/установки", "Для чего нужна"],
        [
            ["Node.js", "20+ / использовался Node.js 22 на VPS", "node --version", "Запуск frontend-сборки и backend scraper"],
            ["npm", "Актуальная версия из Node.js", "npm --version", "Установка зависимостей и запуск npm scripts"],
            ["Vite", "7.x", "npm install; npm run build", "Сборка frontend"],
            ["TypeScript", "5.x", "npm run build", "Компиляция клиентского realtime-слоя"],
            ["SpaceTimeDB CLI", "Совместимый с SDK 2.0", "spacetime --version", "Публикация модуля и генерация bindings"],
            ["SpaceTimeDB SDK", "2.0.0", "npm install", "Клиентское подключение к базе realtime"],
            ["Docker / Docker Compose", "24+ / compose v2", "docker compose version", "Запуск nginx и backend scraper на VPS"],
            ["Nginx", "1.28 в контейнере", "docker compose up -d", "Раздача статики, SSL, proxy для /api"],
        ],
    )

    add_heading(document, "5. Конфигурация проекта")
    add_table(
        document,
        ["Переменная", "Пример значения", "Назначение", "Обязательна?", "Комментарий"],
        [
            ["VITE_SPACETIMEDB_HOST", "https://maincloud.spacetimedb.com", "Адрес SpaceTimeDB", "Да", "Используется frontend realtime-слоем"],
            ["VITE_SPACETIMEDB_DB_NAME", "site-cursors", "Имя базы SpaceTimeDB", "Да", "Должно совпадать с опубликованным модулем"],
            ["HOST", "0.0.0.0", "Адрес прослушивания backend scraper", "Нет", "На VPS нужен 0.0.0.0 внутри контейнера"],
            ["PORT", "18080", "Порт backend scraper", "Да", "Nginx проксирует /api на track-backend:18080"],
            ["TRACK_STORE_PATH", "/app/data/tracks.json", "Путь к JSON-хранилищу заявок", "Нет", "Временное решение до перехода на БД"],
            ["NODE_ENV", "production", "Режим запуска", "Нет", "Для production-среды"],
        ],
    )
    add_paragraph(
        document,
        "Пример .env.example для проекта:"
    )
    add_paragraph(
        document,
        "VITE_SPACETIMEDB_HOST=https://maincloud.spacetimedb.com\nVITE_SPACETIMEDB_DB_NAME=site-cursors\nHOST=0.0.0.0\nPORT=18080\nTRACK_STORE_PATH=./backend_scraper/data/tracks.json\nNODE_ENV=production"
    )

    add_heading(document, "6. Инструкция запуска")
    add_table(
        document,
        ["№", "Действие", "Команда", "Ожидаемый результат", "Скриншот"],
        [
            ["1", "Скачать проект", "git clone https://github.com/Desai0/site", "Папка проекта появилась локально", "01_repo_structure.png"],
            ["2", "Установить frontend-зависимости", "npm install", "Зависимости установлены без критических ошибок", "02_install_dependencies.png"],
            ["3", "Создать настройки", "cp .env.example .env.production", "Файл настроек создан без секретов", "03_env_example.png"],
            ["4", "Собрать frontend", "npm run build", "В папке dist появилась production-сборка", "04_success_run.png"],
            ["5", "Запустить backend scraper", "cd backend_scraper && HOST=0.0.0.0 PORT=18080 npm start", "Backend слушает порт 18080", "04_success_run.png"],
            ["6", "Проверить backend", "curl http://127.0.0.1:18080/health", "Ответ: {\"ok\":true}", "logs/app_start_log.txt"],
            ["7", "Проверить заявку", "curl -X POST /api/tracks/submit ...", "Сервер возвращает id и downloadUrl/status", "05_app_working.png"],
            ["8", "Открыть сайт", "http://127.0.0.1:5173 или production-домен", "Интерфейс сайта доступен", "05_app_working.png"],
        ],
    )

    add_heading(document, "7. Скриншоты проверки")
    add_paragraph(
        document,
        "Для сдачи этапа необходимо приложить скриншоты в папку screenshots/. В отчёте должны быть ссылки или вставки следующих изображений:"
    )
    add_bullets(
        document,
        [
            "01_repo_structure.png — структура проекта в VS Code или GitHub.",
            "02_install_dependencies.png — выполнение npm install без критических ошибок.",
            "03_env_example.png — пример .env.example без настоящих секретов.",
            "04_success_run.png — успешный запуск frontend/backend.",
            "05_app_working.png — открытая страница сайта или результат запроса к API.",
            "06_git_status.png — состояние git после подготовки файлов.",
        ],
    )

    add_heading(document, "8. Журнал найденных проблем")
    add_table(
        document,
        ["ID", "Проблема", "Как проявляется", "Причина", "Что сделать", "Статус"],
        [
            ["P-01", "Не хватало актуальной инструкции запуска", "Другой пользователь не понимает порядок запуска frontend, SpaceTimeDB и scraper", "Проект развивался поэтапно", "Обновить README.md и DEPLOY_VPS.md", "В работе"],
            ["P-02", "Backend scraper конфликтовал с портом 8080", "Ошибка EADDRINUSE при npm start", "Порт был занят другим сервисом", "Перенести backend на PORT=18080", "Исправлено"],
            ["P-03", "Nginx не видел backend на host.docker.internal", "504 Gateway Timeout при запросе /api", "Контейнер nginx не мог подключиться к host-gateway", "Запустить backend как сервис track-backend в одной docker-сети", "Исправлено"],
            ["P-04", "История рисунков падала при live-подписке на таблицу", "Ошибка columns undefined в SDK", "Проблема совместимости client SDK с таблицей истории", "Получать историю через procedure getStrokeHistory без live subscription", "Исправлено"],
            ["P-05", "Хранилище заявок временное", "Данные лежат в JSON/локальном файле", "Быстрая реализация прототипа", "На следующем этапе перейти на полноценную БД, рассмотреть PostgreSQL/Redis", "Запланировано"],
            ["P-06", "Scraper зависит от внешних сайтов", "Поиск может сломаться при изменении верстки или блокировке", "HTML парсится с внешних ресурсов", "Добавить обработку ошибок, кэширование и альтернативные провайдеры", "Запланировано"],
        ],
    )

    add_heading(document, "9. Вывод о готовности проекта")
    add_paragraph(
        document,
        "По итогам входного аудита проект частично готов к дальнейшей эксплуатации и сопровождению. Основной frontend запускается и собирается через Vite, realtime-функции вынесены в SpaceTimeDB, а страница заявок получила backend-прототип для поиска и сохранения ссылок скачивания треков. Также определены зависимости, переменные окружения и порядок запуска на локальной машине и VPS."
    )
    add_paragraph(
        document,
        "Проект сохраняет связь с работой УП.02: текущий этап не является созданием новой системы, а продолжает развитие уже начатого проекта. На УП.03 акцент смещён на сопровождение, проверку состава, настройку окружения, документирование и подготовку к дальнейшему развитию серверной части."
    )
    add_paragraph(
        document,
        "К следующему этапу проект требует доработки документации, подготовки .env.example, оформления скриншотов и улучшения backend-хранилища. В дальнейшем планируется перенести scraper во взаимодействие с отдельным сервером БД, заменить временное локальное хранение на более надёжное решение и рассмотреть Redis для кэширования результатов поиска."
    )
    add_paragraph(document, "Основные задачи на следующий этап:")
    add_bullets(
        document,
        [
            "Актуализировать README.md, DEPLOY_VPS.md и инструкцию запуска.",
            "Подготовить .env.example и таблицу конфигурации без секретов.",
            "Перенести хранение заявок из временного JSON/SQLite-подхода в полноценную БД.",
            "Добавить обработку ошибок и журналирование для backend scraper.",
            "Подготовить скриншоты запуска, работы сайта и проверки API.",
        ],
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    build()
