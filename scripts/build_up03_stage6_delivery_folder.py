from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "УП03_Этап6"
DOCS = OUT / "docs"
REPORTS = OUT / "reports"
SCREENSHOTS = OUT / "screenshots"
PROJECT = OUT / "project_files_in_repository"


def write(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + "\n", encoding="utf-8")


def copy(source: Path, destination: Path):
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)


def set_cell_text(cell, value, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        table.rows[0].cells[index].width = Inches(widths[index])
        set_cell_text(table.rows[0].cells[index], header, True)
        shade_cell(table.rows[0].cells[index], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = Inches(widths[index])
            set_cell_text(cells[index], value)
    document.add_paragraph()


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(46, 116, 181)


def add_paragraph(document, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        first, rest = text.split(":", 1)
        paragraph.add_run(first + ":").bold = True
        paragraph.add_run(rest)
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)


def create_support_report():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("ОТЧЁТ ПО СОПРОВОЖДЕНИЮ ПО")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("УП.03. Этап 6. CI/CD, релиз и инцидент поддержки")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(80, 80, 80)

    add_table(
        document,
        ["Параметр", "Значение"],
        [
            ["Проект", "Desaichk Portfolio и backend заявок на музыкальные треки"],
            ["Репозиторий", "https://github.com/Desai0/site"],
            ["Рабочая ветка", "support-fix-cors-policy"],
            ["GitHub Issue", "https://github.com/Desai0/site/issues/9"],
            ["Pull Request", "https://github.com/Desai0/site/pull/11"],
            ["Версия релиза", "0.3.1"],
            ["Дата проверки", "10 июня 2026 года"],
        ],
        [1.7, 4.9],
    )

    add_heading(document, "1. Выбранная проблема")
    add_paragraph(
        document,
        "В этапе 5 был обнаружен риск SEC-02: backend отвечал заголовком "
        "Access-Control-Allow-Origin: * на запросы с любого сайта. Это позволяло "
        "стороннему origin читать ответы публичного API из браузера пользователя."
    )

    add_heading(document, "2. Воспроизведение")
    for text in [
        "Запустить исходную версию backend или проверить production endpoint.",
        "Отправить OPTIONS-запрос с Origin: https://evil.example.",
        "Убедиться, что ответ содержит Access-Control-Allow-Origin: *.",
    ]:
        add_bullet(document, text)

    add_heading(document, "3. Диагностика")
    add_paragraph(
        document,
        "Проверка backend_scraper/index.js показала, что функция setCorsHeaders "
        "безусловно устанавливала wildcard. Конфигурация CORS_ORIGINS уже была "
        "описана в production env-примере, но код backend её не использовал."
    )
    add_paragraph(
        document,
        "Фактические команды и вывод сохранены в reports/03_bug_reproduced.txt "
        "и reports/04_logs_diagnostics.txt."
    )

    add_heading(document, "4. Исправление")
    for text in [
        "Добавлен разбор списка CORS_ORIGINS через запятую.",
        "Access-Control-Allow-Origin возвращается только для разрешённого origin.",
        "Для разрешённого origin добавляется Vary: Origin.",
        "Docker Compose передаёт production и local значения CORS_ORIGINS.",
        "Добавлен интеграционный тест backend_scraper/security_test.js.",
    ]:
        add_bullet(document, text)

    add_heading(document, "5. Проверка")
    add_table(
        document,
        ["Команда", "Назначение", "Ожидаемый результат"],
        [
            ["npm run check:backend", "Проверка синтаксиса backend", "Exit code 0"],
            ["npm run test:security", "Allowed/blocked origin", "Три строки PASS"],
            ["npm run build", "Production-сборка Vite", "Build completed"],
            ["npm run release:check", "Полная release-проверка", "Все шаги пройдены"],
        ],
        [2.0, 2.4, 2.2],
    )

    add_heading(document, "6. CI/CD")
    add_paragraph(
        document,
        "Workflow .github/workflows/ci.yml запускается на push и pull_request. "
        "Он использует Node.js 22, выполняет npm ci и npm run check. GitHub "
        "Actions для Pull Request #11 завершён со статусом success."
    )

    add_heading(document, "7. Релиз")
    add_paragraph(
        document,
        "Подготовлена версия 0.3.1. Обновлены CHANGELOG.md, RELEASE_NOTES.md и "
        "RELEASE_CHECKLIST.md. Скрипт scripts/create-release.bat выполняет "
        "release-check и создаёт архив release/site-0.3.1.zip."
    )

    add_heading(document, "8. Итог")
    add_paragraph(
        document,
        "Причина инцидента найдена и устранена. Создан Issue #9, исправление "
        "выполнено в отдельной ветке и отправлено в Pull Request #11. Локальные "
        "проверки и GitHub Actions прошли успешно. Тег v0.3.1 создаётся после "
        "принятия Pull Request; для сдачи уже подготовлен release-архив."
    )

    add_heading(document, "9. Скриншоты")
    screenshots = [
        "01_github_issue_created.png",
        "02_branch_created.png",
        "03_bug_reproduced.png",
        "04_logs_diagnostics.png",
        "05_fix_commit.png",
        "06_tests_passed_locally.png",
        "07_github_actions_success.png",
        "08_pull_request.png",
        "09_changelog_release_notes.png",
        "10_release_tag_or_archive.png",
    ]
    for name in screenshots:
        add_bullet(document, name)

    document.save(DOCS / "SUPPORT_REPORT.docx")


def create_markdown_docs():
    write(
        DOCS / "INCIDENT_REPORT.md",
        """
# INCIDENT_REPORT.md

## 1. Инцидент

Backend заявок разрешает CORS-запросы с произвольных сайтов.

## 2. Где обнаружено

Риск SEC-02 обнаружен при проверке production API и исходного кода в этапе 5.

## 3. Как воспроизвести

1. Отправить `OPTIONS /api/tracks/1`.
2. Указать `Origin: https://evil.example`.
3. В исходной версии получить `Access-Control-Allow-Origin: *`.

## 4. Диагностика

В `backend_scraper/index.js` функция `setCorsHeaders` безусловно устанавливала
wildcard. Переменная `CORS_ORIGINS` была описана в env-примере, но backend её не
использовал.

## 5. Причина

Политика CORS была захардкожена и не зависела от окружения или заголовка Origin.

## 6. Исправление

- backend читает список `CORS_ORIGINS`;
- заголовок возвращается только для origin из списка;
- Docker Compose передаёт production/local конфигурацию;
- добавлен интеграционный тест разрешённого и запрещённого origin.

## 7. Проверка

```text
npm run check:backend
npm run test:security
npm run build
npm run release:check
```

Фактический вывод команд находится в папке `reports`.

## 8. Итог

Локальная проверка подтверждает исправление. Созданы GitHub Issue #9 и Pull
Request #11. GitHub Actions CI завершён успешно. Тег создаётся после
принятия Pull Request; локальный release-архив уже подготовлен.
""",
    )
    write(
        DOCS / "ISSUE_BODY.md",
        """
# Проблема: backend разрешает CORS для любого сайта

## Описание

API заявок возвращает `Access-Control-Allow-Origin: *`. Сторонняя веб-страница
может читать ответы API из браузера пользователя.

## Шаги воспроизведения

1. Отправить OPTIONS-запрос к `/api/tracks/1`.
2. Добавить `Origin: https://evil.example`.
3. Проверить заголовки ответа.

## Ожидаемый результат

Посторонний origin не получает `Access-Control-Allow-Origin`.

## Фактический результат

Исходная версия возвращает `Access-Control-Allow-Origin: *`.

## Окружение

- Windows 10/11 и Debian 12
- Node.js 22
- local/Docker/production

## Доказательства

`УП03_Этап6/reports/03_bug_reproduced.txt`
""",
    )
    write(
        DOCS / "PULL_REQUEST_BODY.md",
        """
# Что изменено

- CORS ограничен списком `CORS_ORIGINS`.
- Добавлен тест разрешённого и запрещённого origin.
- Добавлен GitHub Actions workflow.
- Обновлены release-документы и команды проверки.

# Связанный issue

Closes #ISSUE_NUMBER

# Как проверено

- `npm run check:backend`
- `npm run test:security`
- `npm run build`
- `npm run release:check`

# Скриншоты

Добавлены в `УП03_Этап6/screenshots`.
""",
    )


def copy_project_files():
    files = [
        (ROOT / ".github" / "workflows" / "ci.yml", PROJECT / ".github" / "workflows" / "ci.yml"),
        (ROOT / "CHANGELOG.md", PROJECT / "CHANGELOG.md"),
        (ROOT / "RELEASE_NOTES.md", PROJECT / "RELEASE_NOTES.md"),
        (ROOT / "RELEASE_CHECKLIST.md", PROJECT / "RELEASE_CHECKLIST.md"),
        (ROOT / "README.md", PROJECT / "README.md"),
        (ROOT / "DEPLOYMENT.md", PROJECT / "DEPLOYMENT.md"),
        (ROOT / "scripts" / "test.bat", PROJECT / "scripts" / "test.bat"),
        (ROOT / "scripts" / "build.bat", PROJECT / "scripts" / "build.bat"),
        (ROOT / "scripts" / "release-check.bat", PROJECT / "scripts" / "release-check.bat"),
        (ROOT / "scripts" / "create-release.bat", PROJECT / "scripts" / "create-release.bat"),
        (ROOT / "backend_scraper" / "security_test.js", PROJECT / "backend_scraper" / "security_test.js"),
    ]
    for source, destination in files:
        copy(source, destination)


def main():
    for directory in [DOCS, REPORTS, SCREENSHOTS, PROJECT]:
        directory.mkdir(parents=True, exist_ok=True)

    write(OUT / "repo_link.txt", "https://github.com/Desai0/site")
    create_markdown_docs()
    copy(ROOT / "RELEASE_NOTES.md", DOCS / "RELEASE_NOTES.md")
    copy(ROOT / "RELEASE_CHECKLIST.md", DOCS / "RELEASE_CHECKLIST.md")
    copy(ROOT / "CHANGELOG.md", DOCS / "CHANGELOG.md")
    copy_project_files()
    create_support_report()
    print(f"Created: {OUT}")


if __name__ == "__main__":
    main()
